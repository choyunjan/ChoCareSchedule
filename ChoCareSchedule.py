import calendar
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from ortools.sat.python import cp_model
import pandas as pd
import streamlit as st

st.set_page_config(page_title="護理之家自動排班系統", layout="wide")
st.title("🏥 卓醫院附設護理之家 - 自動排班系統")

# ==========================================
# 1. 預設實際員工名單 (2樓 17人 / 6樓 8人)
# ==========================================
DEFAULT_EMP_2F = [
    "蕭切",
    "廖淑華",
    "鄭鳳英",
    "劉春燕",
    "梅氏尚",
    "范氏映",
    "黃氏梅",
    "陳百毓",
    "陳氏蘭",
    "徐曉涵",
    "鄭銘傑",
    "吳秉儒",
    "農氏深",
    "林氏專", "許氏緬",
    "杜氏享",
    "陳美雯",
]

DEFAULT_EMP_6F = [
    "郭秋貴",
    "陳氏幸",
    "杜氏深",
    "阮氏銀",
    "陳氏恆",
    "洪紹棠",
    "盧氏念",
    "蔡雪晏",
]

# ==========================================
# 側邊欄：控制參數設定
# ==========================================
st.sidebar.header("⚙️ 排班參數設定")

# 1. 年月設定
col1, col2 = st.sidebar.columns(2)
with col1:
    YEAR = st.number_input("年份", min_value=2024, max_value=2030, value=2026)
with col2:
    MONTH = st.number_input("月份", min_value=1, max_value=12, value=9)

DAYS_IN_MONTH = calendar.monthrange(YEAR, MONTH)[1]

# 2. 國定假日設定
holidays_input = st.sidebar.text_input(
    "當月國定假日日期 (以逗號分隔)",
    value="",
    help="例如：9月28日放假，請輸入 28",
)
HOLIDAYS_LIST = [
    int(x.strip())
    for x in holidays_input.split(",")
    if x.strip().isdigit() and 1 <= int(x.strip()) <= DAYS_IN_MONTH
]

# 計算當月應休天數基準
WEEKDAY_NAMES = ["一", "二", "三", "四", "五", "六", "日"]
target_off_days = sum(
    1
    for d in range(1, DAYS_IN_MONTH + 1)
    if date(YEAR, MONTH, d).weekday() in [5, 6] or d in HOLIDAYS_LIST
)

st.sidebar.info(
    f"📅 **{YEAR}年{MONTH}月** 共 {DAYS_IN_MONTH} 天\n\n當月例假日+國定假日基準：**{target_off_days} 天**"
)

# 3. 人員名單設定 (可在介面上直接微調或增加同仁)
st.sidebar.subheader("👥 人員名單微調")
emp_2f_text = st.sidebar.text_area(
    "2樓人員名單 (每行或逗號分隔)",
    value="\n".join(DEFAULT_EMP_2F),
    height=150,
)
emp_6f_text = st.sidebar.text_area(
    "6樓人員名單 (每行或逗號分隔)",
    value="\n".join(DEFAULT_EMP_6F),
    height=120,
)

EMP_2F = [
    x.strip()
    for x in emp_2f_text.replace(",", "\n").split("\n")
    if x.strip()
]
EMP_6F = [
    x.strip()
    for x in emp_6f_text.replace(",", "\n").split("\n")
    if x.strip()
]
ALL_EMP = EMP_2F + EMP_6F

# 4. 純白班同仁設定
st.sidebar.subheader("☀️ 特定班別限制")
fixed_day_input = st.sidebar.text_area(
    "固定只排白班的人員",
    value="蕭切\n郭秋貴",
    help="請輸入同仁姓名，需與上述名單一致",
)
FIXED_DAY_WORKERS = [
    x.strip()
    for x in fixed_day_input.replace(",", "\n").split("\n")
    if x.strip()
]

# ==========================================
# 主畫面：啟動排班運算
# ==========================================
if st.button("🚀 開始產生自動排班表", type="primary"):
    with st.spinner("AI 求解器正在計算最優班表，請稍候..."):

        ALL_SHIFTS = ["D2", "E2", "N2", "D6", "E6", "N6", "R"]
        DAILY_REQUIREMENTS = {
            "D2": 4,
            "E2": 3,
            "N2": 3,
            "D6": 2,
            "E6": 2,
            "N6": 2,
        }

        model = cp_model.CpModel()
        shifts = {}

        for e in ALL_EMP:
            for d in range(1, DAYS_IN_MONTH + 1):
                for s in ALL_SHIFTS:
                    shifts[(e, d, s)] = model.NewBoolVar(f"shift_{e}_{d}_{s}")

        # 約束 1: 每人每天一班
        for e in ALL_EMP:
            for d in range(1, DAYS_IN_MONTH + 1):
                model.Add(sum(shifts[(e, d, s)] for s in ALL_SHIFTS) == 1)

        # 約束 2: 每日人力需求
        for d in range(1, DAYS_IN_MONTH + 1):
            for s, count in DAILY_REQUIREMENTS.items():
                model.Add(sum(shifts[(e, d, s)] for e in ALL_EMP) == count)

        # 約束 3: 休假數控制
        for e in ALL_EMP:
            model.Add(
                sum(shifts[(e, d, "R")] for d in range(1, DAYS_IN_MONTH + 1))
                >= target_off_days
            )
            model.Add(
                sum(shifts[(e, d, "R")] for d in range(1, DAYS_IN_MONTH + 1))
                <= target_off_days + 4
            )

        # 約束 4: 連續工作不超過 6 天
        for e in ALL_EMP:
            for d in range(1, DAYS_IN_MONTH - 5):
                model.Add(sum(shifts[(e, d + i, "R")] for i in range(7)) >= 1)

        # 約束 5: 轉班限制 (>12小時)
        for e in ALL_EMP:
            for d in range(1, DAYS_IN_MONTH):
                e_today = [shifts[(e, d, "E2")], shifts[(e, d, "E6")]]
                d_next = [shifts[(e, d + 1, "D2")], shifts[(e, d + 1, "D6")]]
                for es in e_today:
                    for ds in d_next:
                        model.AddImplication(es, ds.Not())

                n_today = [shifts[(e, d, "N2")], shifts[(e, d, "N6")]]
                de_next = (
                    d_next
                    + [shifts[(e, d + 1, "E2")], shifts[(e, d + 1, "E6")]]
                )
                for ns in n_today:
                    for des in de_next:
                        model.AddImplication(ns, des.Not())

        # 約束 6: 固定純白班
        for e in FIXED_DAY_WORKERS:
            if e in ALL_EMP:
                for d in range(1, DAYS_IN_MONTH + 1):
                    model.Add(shifts[(e, d, "E2")] == 0)
                    model.Add(shifts[(e, d, "N2")] == 0)
                    model.Add(shifts[(e, d, "E6")] == 0)
                    model.Add(shifts[(e, d, "N6")] == 0)

        # 目標函數: 跨樓層懲罰
        penalty_list = []
        for e in EMP_2F:
            for d in range(1, DAYS_IN_MONTH + 1):
                for s in ["D6", "E6", "N6"]:
                    penalty_list.append(shifts[(e, d, s)] * 10)

        for e in EMP_6F:
            for d in range(1, DAYS_IN_MONTH + 1):
                for s in ["D2", "E2", "N2"]:
                    penalty_list.append(shifts[(e, d, s)] * 10)

        model.Minimize(sum(penalty_list))

        solver = cp_model.CpSolver()
        solver.parameters.max_time_in_seconds = 15.0
        status = solver.Solve(model)

    if status in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        st.success(
            f"🎉 成功找到最優排班方案！全院跨樓層支援僅 {int(solver.ObjectiveValue()//10)} 人次。"
        )

        wb = Workbook()
        ws = wb.active
        ws.title = f"{MONTH}月班表"

        red_fill = PatternFill(
            start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"
        )
        yellow_fill = PatternFill(
            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
        )
        floor2_fill = PatternFill(
            start_color="DCE6F1", end_color="DCE6F1", fill_type="solid"
        )
        floor6_fill = PatternFill(
            start_color="F2DCD1", end_color="F2DCD1", fill_type="solid"
        )
        header_fill = PatternFill(
            start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
        )

        align_center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3"),
        )

        # 表頭繪製
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=DAYS_IN_MONTH + 3,
        )
        title_cell = ws.cell(
            row=1,
            column=1,
            value=f"卓醫院附設護理之家 {YEAR}年度{MONTH}月照顧服務員班表",
        )
        title_cell.font = Font(size=16, bold=True)
        title_cell.alignment = align_center

        ws.cell(row=2, column=1, value="原樓層").alignment = align_center
        ws.cell(row=2, column=2, value="姓名").alignment = align_center
        ws.merge_cells("A2:A3")
        ws.merge_cells("B2:B3")

        for d in range(1, DAYS_IN_MONTH + 1):
            dt = date(YEAR, MONTH, d)
            col_idx = d + 2
            cell_day = ws.cell(row=2, column=col_idx, value=d)
            cell_day.alignment = align_center
            cell_day.font = Font(bold=True)
            cell_day.fill = header_fill

            weekday_str = WEEKDAY_NAMES[dt.weekday()]
            cell_week = ws.cell(row=3, column=col_idx, value=weekday_str)
            cell_week.alignment = align_center

            if dt.weekday() in [5, 6] or d in HOLIDAYS_LIST:
                cell_week.fill = red_fill
                cell_week.font = Font(color="9C0006", bold=True)
            else:
                cell_week.fill = header_fill

        total_off_col = DAYS_IN_MONTH + 3
        ws.cell(row=2, column=total_off_col, value="總休天數").alignment = (
            align_center
        )
        ws.merge_cells(
            start_row=2,
            start_column=total_off_col,
            end_row=3,
            end_column=total_off_col,
        )

        # 填入排班資料
        schedule_data = []
        for idx, e in enumerate(ALL_EMP):
            curr_row = 4 + idx
            orig_floor = "2樓" if e in EMP_2F else "6樓"
            ws.cell(row=curr_row, column=1, value=orig_floor).alignment = (
                align_center
            )
            ws.cell(row=curr_row, column=2, value=e).alignment = align_center

            emp_row_data = {"原樓層": orig_floor, "姓名": e}
            off_count = 0
            for d in range(1, DAYS_IN_MONTH + 1):
                col_idx = d + 2
                for s in ALL_SHIFTS:
                    if solver.Value(shifts[(e, d, s)]) == 1:
                        cell = ws.cell(row=curr_row, column=col_idx, value=s)
                        cell.alignment = align_center
                        emp_row_data[f"{d}日"] = s

                        if s == "R":
                            cell.fill = yellow_fill
                            off_count += 1
                        elif s in ["D2", "E2", "N2"]:
                            cell.fill = floor2_fill
                        elif s in ["D6", "E6", "N6"]:
                            cell.fill = floor6_fill

            ws.cell(row=curr_row, column=total_off_col, value=off_count).alignment = align_center
            emp_row_data["總休天數"] = off_count
            schedule_data.append(emp_row_data)

        for r in range(1, 4 + len(ALL_EMP)):
            for c in range(1, total_off_col + 1):
                ws.cell(row=r, column=c).border = thin_border

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 14
        for c in range(3, total_off_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = 4.5

        # 網頁預覽與下載
        df_show = pd.DataFrame(schedule_data)
        st.dataframe(df_show, use_container_width=True)

        excel_buffer = BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        st.download_button(
            label="📥 下載完整 Excel 排班表",
            data=excel_buffer,
            file_name=f"{YEAR}年{MONTH}月_照顧服務員班表.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        st.error("❌ 排班失敗：目前的條件過於嚴格或人力不足，請調整側邊欄的限制。")